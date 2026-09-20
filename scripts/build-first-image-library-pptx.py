import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


NAVY = RGBColor(0x0B, 0x14, 0x26)
CORAL = RGBColor(0xF0, 0x6E, 0x5F)
IVORY = RGBColor(0xF7, 0xF4, 0xEE)
IVORY_DARK = RGBColor(0xE9, 0xE4, 0xDA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x67, 0x70, 0x85)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    width: float,
    height: float,
    *,
    size: float,
    color: RGBColor = NAVY,
    bold: bool = False,
    font: str = "Aptos",
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
):
    box = slide.shapes.add_textbox(
        Inches(x), Inches(y), Inches(width), Inches(height)
    )
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_contained_picture(slide, image_path: Path, box):
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    x, y, width, height = box
    scale = min(width / image_width, height / image_height)
    rendered_width = image_width * scale
    rendered_height = image_height * scale
    return slide.shapes.add_picture(
        str(image_path),
        Inches(x + (width - rendered_width) / 2),
        Inches(y + (height - rendered_height) / 2),
        Inches(rendered_width),
        Inches(rendered_height),
    )


def build_deck(manifest_path: Path, output_path: Path) -> None:
    manifest = json.loads(manifest_path.read_text())
    entries = manifest["entries"]
    section_title = manifest["sectionTitle"]
    resources = manifest["resources"]
    project_root = manifest_path.resolve().parents[2]

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    presentation.core_properties.author = "ASOPRS Study Portal"
    presentation.core_properties.title = f"ASOPRS Image Library — {section_title}"
    presentation.core_properties.subject = f"{section_title} image library"
    blank_layout = presentation.slide_layouts[6]

    title_slide = presentation.slides.add_slide(blank_layout)
    background = title_slide.background.fill
    background.solid()
    background.fore_color.rgb = NAVY
    accent = title_slide.shapes.add_shape(
        MSO_AUTO_SHAPE_TYPE.RECTANGLE, 0, 0, Inches(0.16), Inches(7.5)
    )
    accent.fill.solid()
    accent.fill.fore_color.rgb = CORAL
    accent.line.fill.background()
    add_text(
        title_slide,
        "ASOPRS IMAGE LIBRARY",
        0.85,
        1.25,
        7.2,
        0.35,
        size=15,
        color=CORAL,
        bold=True,
    )
    add_text(
        title_slide,
        section_title,
        0.85,
        1.75,
        10.8,
        0.85,
        size=42,
        color=WHITE,
        bold=True,
        font="Aptos Display",
    )
    add_text(
        title_slide,
        f"{len(resources)} resources · {len(entries)} figures · curriculum order",
        0.85,
        2.82,
        8.5,
        0.5,
        size=20,
        color=RGBColor(0xD9, 0xDE, 0xE8),
    )
    add_text(
        title_slide,
        "Review draft · Section 1",
        0.85,
        6.55,
        6,
        0.3,
        size=13,
        color=RGBColor(0x9C, 0xA8, 0xBD),
    )

    for entry in entries:
        slide = presentation.slides.add_slide(blank_layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = WHITE
        add_text(
            slide,
            entry["documentTitle"],
            0.7,
            0.35,
            9.5,
            0.55,
            size=24,
            bold=True,
            font="Aptos Display",
            valign=MSO_ANCHOR.MIDDLE,
        )
        add_text(
            slide,
            f'{entry["figureLabel"]} · Page {entry["pageNumber"]}',
            10.2,
            0.45,
            2.4,
            0.3,
            size=12,
            color=CORAL,
            bold=True,
            align=PP_ALIGN.RIGHT,
        )
        image_panel = slide.shapes.add_shape(
            MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE,
            Inches(0.7),
            Inches(1.1),
            Inches(7.5),
            Inches(5.65),
        )
        image_panel.fill.solid()
        image_panel.fill.fore_color.rgb = IVORY
        image_panel.line.color.rgb = IVORY_DARK
        image_path = project_root / "public" / entry["imagePath"].lstrip("/")
        add_contained_picture(slide, image_path, (0.9, 1.3, 7.1, 5.25))
        add_text(
            slide,
            "DESCRIPTION",
            8.65,
            1.3,
            3.8,
            0.3,
            size=12,
            color=CORAL,
            bold=True,
        )
        caption_size = 18 if len(entry["caption"]) < 220 else 15
        add_text(
            slide,
            entry["caption"],
            8.65,
            1.77,
            3.75,
            3.85,
            size=caption_size,
        )
        add_text(
            slide,
            f'Source: {entry["sourcePdfPath"]}',
            8.65,
            6.25,
            3.75,
            0.45,
            size=10,
            color=MUTED,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output_path)
    print(f"Wrote {len(presentation.slides)} slides to {output_path}")


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit(
            "usage: build-first-image-library-pptx.py <manifest.json> <output.pptx>"
        )
    build_deck(Path(sys.argv[1]), Path(sys.argv[2]))


if __name__ == "__main__":
    main()
