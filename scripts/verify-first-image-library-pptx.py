from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


DECK_PATH = (
    Path(__file__).resolve().parents[1]
    / "public"
    / "image-library"
    / "asoprs-image-library-acquired-laxity.pptx"
)
EXPECTED_IMAGE_SLIDES = 24


def main() -> None:
    presentation = Presentation(DECK_PATH)
    assert len(presentation.slides) == EXPECTED_IMAGE_SLIDES + 1, (
        f"expected one title slide plus {EXPECTED_IMAGE_SLIDES} image slides, "
        f"found {len(presentation.slides)} slides"
    )

    for slide_number, slide in enumerate(list(presentation.slides)[1:], start=2):
        pictures = [
            shape for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        ]
        assert len(pictures) == 1, (
            f"slide {slide_number} must contain exactly one embedded image; "
            f"found {len(pictures)}"
        )
        picture = pictures[0]
        assert picture.width > 0 and picture.height > 0, (
            f"slide {slide_number} has a zero-sized image"
        )
        assert picture.image.blob, f"slide {slide_number} image data is empty"

    print(
        f"Verified {len(presentation.slides)} slides with "
        f"{EXPECTED_IMAGE_SLIDES} embedded images: {DECK_PATH}"
    )


if __name__ == "__main__":
    main()
